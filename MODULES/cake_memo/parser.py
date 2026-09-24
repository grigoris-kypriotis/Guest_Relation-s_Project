"""
Document parsing for Cake Memo entries.

Parses a cake memo Word document by opening it, mapping header cells to column indices
(header-driven, never hardcoded), and reading row 1 using field_format.py's parse functions.
Symmetrical counterpart to generate_cake_memo_document.
"""

import os
import re
from docx import Document

from MODULES.cake_memo.paths import CAKE_MEMO_TEMPLATE_PATH
from MODULES.cake_memo.field_format import (
    parse_service_description,
    parse_provided_at,
    parse_charge,
    FLAVOR_DISPLAY,
)


def parse_cake_memo_document(docx_path: str) -> dict:
    """
    Parse a cake memo Word document into structured form data.

    Mirrors the header-driven pattern from generate_cake_memo_document.
    Opens the file, finds the header row, maps columns by name (case-insensitive),
    reads row 1 only (confirmed: one memo = one room = one document).

    Returns a dict with the same keys as generate_cake_memo_document's form_data input,
    so that form.set_form_data(parse_cake_memo_document(path)) works directly.

    Args:
        docx_path: path to a cake memo docx file

    Returns:
        dict with keys:
            - flavor: str (one of FLAVOR_DISPLAY keys)
            - written_text: str
            - qty: int (defaults to 1 if not parseable as int)
            - venue: str (one of VENUE_DISPLAY keys)
            - hour: int
            - minute: int
            - is_pm: bool
            - charge_state: str (one of CHARGE_*)
            - complimentary_by: str
            - room_number: str

        Note: pax is omitted from the returned dict (it's a form-only field,
        never written to or read from the docx). Caller can set pax to a default if needed.

    Raises:
        FileNotFoundError: if docx_path does not exist
        ValueError: if required columns are missing, or if data in critical columns
                   is malformed and cannot be parsed
    """
    docx_path = os.path.abspath(docx_path)

    if not os.path.exists(docx_path):
        raise FileNotFoundError(f"Document not found: {docx_path}")

    # Open document
    doc = Document(docx_path)

    if not doc.tables:
        raise ValueError("Document has no tables")

    table = doc.tables[0]

    # Map header row (row 0) to column indices
    col_indices = {}
    expected_headers = {
        "SERVICE DESCRIPTION": None,
        "QTY": None,
        "PROVIDED AT": None,
        "DATE": None,
        "CHARGE": None,
        "ROOM NUMBER": None,
    }

    for idx, cell in enumerate(table.rows[0].cells):
        # Normalize: remove line breaks, tabs, etc., then strip
        header_text = re.sub(r'[\r\n\t\a]', ' ', cell.text).strip()

        # Match against expected headers (case-insensitive)
        for expected in expected_headers.keys():
            if re.search(f"(?i)^{re.escape(expected)}$", header_text):
                col_indices[expected] = idx
                break

    # Check all required columns are found
    missing = [name for name in expected_headers.keys() if name not in col_indices]
    if missing:
        raise ValueError(
            f"Required header columns missing in document: {', '.join(missing)}"
        )

    # Read row 1 only
    if len(table.rows) < 2:
        raise ValueError("Table must have at least 2 rows (header + data)")

    row = table.rows[1]

    # Extract raw cell texts
    service_desc_text = row.cells[col_indices["SERVICE DESCRIPTION"]].text.strip()
    qty_text = row.cells[col_indices["QTY"]].text.strip()
    provided_at_text = row.cells[col_indices["PROVIDED AT"]].text.strip()
    # date_text = row.cells[col_indices["DATE"]].text.strip()  # Not round-tripped
    charge_text = row.cells[col_indices["CHARGE"]].text.strip()
    room_number_text = row.cells[col_indices["ROOM NUMBER"]].text.strip()

    # Parse SERVICE DESCRIPTION
    flavor, written_text = parse_service_description(service_desc_text)
    if flavor is None:
        # Malformed SERVICE DESCRIPTION; raise rather than silently default
        raise ValueError(
            f"Cannot parse SERVICE DESCRIPTION: '{service_desc_text}' "
            f"does not match expected format"
        )

    # Parse QTY (defensively: default to 1 if not parseable as int)
    try:
        qty = int(qty_text)
    except (ValueError, TypeError):
        qty = 1

    # Parse PROVIDED AT
    venue, hour, minute, is_pm = parse_provided_at(provided_at_text)
    if venue is None or hour is None or minute is None or is_pm is None:
        # Malformed PROVIDED AT; raise rather than silently default
        raise ValueError(
            f"Cannot parse PROVIDED AT: '{provided_at_text}' "
            f"does not match expected format (venue time_pattern)"
        )

    # Parse CHARGE (always safe — parse_charge defaults to CHARGE_PAID)
    charge_state, complimentary_by = parse_charge(charge_text)

    # ROOM NUMBER is as-is string
    room_number = room_number_text

    return {
        "flavor": flavor,
        "written_text": written_text,
        "qty": qty,
        "venue": venue,
        "hour": hour,
        "minute": minute,
        "is_pm": is_pm,
        "charge_state": charge_state,
        "complimentary_by": complimentary_by,
        "room_number": room_number,
    }
