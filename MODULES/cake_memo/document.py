"""
Document generation for Cake Memo entries.

Generates a single-row cake memo document by opening the template,
mapping header cells to column indices (header-driven, never hardcoded),
and writing form data into row 1 using the field_format.py compose functions.
"""

import os
import re
from datetime import datetime
from docx import Document

from MODULES.cake_memo.paths import CAKE_MEMO_TEMPLATE_PATH
from MODULES.cake_memo.field_format import (
    compose_service_description,
    compose_provided_at,
    compose_charge,
)


def generate_cake_memo_document(form_data: dict, save_path: str) -> None:
    """
    Generate a cake memo Word document from structured form data.

    Mirrors the header-driven pattern from offers/document.py::generate_word_document.
    Opens the template, finds the header row, maps columns by name (case-insensitive),
    writes only into row 1 (confirmed: one memo = one room = one document).

    Args:
        form_data: dict with keys:
            - flavor: str (one of FLAVOR_DISPLAY keys, e.g., 'strawberry')
            - written_text: str (optional freeform text for cake)
            - qty: int (quantity)
            - venue: str (one of VENUE_DISPLAY keys, e.g., 'il_gusto', 'room')
            - hour: int (0-23)
            - minute: int (0-59)
            - is_pm: bool
            - charge_state: str (one of CHARGE_PAID, CHARGE_PENDING, CHARGE_COMPLIMENTARY)
            - complimentary_by: str (name, if charge_state is CHARGE_COMPLIMENTARY)
            - room_number: str
        save_path: output file path (will be created/overwritten)

    Raises:
        ValueError: if any required header column is missing from the template

    Column mapping:
        - SERVICE DESCRIPTION ← compose_service_description(flavor, written_text)
        - QTY ← str(qty)
        - PROVIDED AT ← compose_provided_at(venue, hour, minute, is_pm)
        - DATE ← today's date in dd/mm format (never from form_data)
        - CHARGE ← compose_charge(charge_state, complimentary_by)
        - ROOM NUMBER ← room_number
    """
    # Normalize path
    save_path = os.path.abspath(save_path)
    template_path = os.path.abspath(CAKE_MEMO_TEMPLATE_PATH)

    # Open template
    doc = Document(template_path)

    if not doc.tables:
        raise ValueError("Template has no tables")

    table = doc.tables[0]

    # Map header row (row 0) to column indices
    # Normalize whitespace in header cells (strip, remove line breaks) for matching
    col_indices = {}
    expected_headers = {
        "SERVICE DESCRIPTION": None,
        "QTY": None,
        "PROVIDED AT": None,
        "DATE": None,
        "CHARGE": None,
        "ROOM NUMBER": None,
    }

    for idx, cell in enumerate(table.rows[0].cells):
        # Normalize: remove line breaks, tabs, etc., then strip
        header_text = re.sub(r'[\r\n\t\a]', ' ', cell.text).strip()

        # Match against expected headers (case-insensitive)
        for expected in expected_headers.keys():
            if re.search(f"(?i)^{re.escape(expected)}$", header_text):
                col_indices[expected] = idx
                break

    # Check all required columns are found
    missing = [name for name in expected_headers.keys() if name not in col_indices]
    if missing:
        raise ValueError(
            f"Required header columns missing in template: {', '.join(missing)}"
        )

    # Write into row 1 only
    if len(table.rows) < 2:
        raise ValueError("Template table must have at least 2 rows (header + data)")

    row = table.rows[1]

    # Compose and write each field
    service_desc = compose_service_description(
        form_data["flavor"], form_data["written_text"]
    )
    row.cells[col_indices["SERVICE DESCRIPTION"]].text = service_desc

    row.cells[col_indices["QTY"]].text = str(form_data["qty"])

    provided_at = compose_provided_at(
        form_data["venue"], form_data["hour"], form_data["minute"], form_data["is_pm"]
    )
    row.cells[col_indices["PROVIDED AT"]].text = provided_at

    # Date: today in dd/mm format
    today_ddmm = datetime.now().strftime("%d/%m")
    row.cells[col_indices["DATE"]].text = today_ddmm

    charge = compose_charge(form_data["charge_state"], form_data["complimentary_by"])
    row.cells[col_indices["CHARGE"]].text = charge

    row.cells[col_indices["ROOM NUMBER"]].text = form_data["room_number"]

    # Save
    doc.save(save_path)
