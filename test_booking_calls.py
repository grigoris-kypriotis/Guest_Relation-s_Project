"""
Unit tests for booking_calls.py:
Tests the 4 scheduling rules, consecutive date resolution, Excel synchronization,
and state persistence.
"""

import os
import shutil
import tempfile
import unittest
from datetime import date, timedelta

import openpyxl

from booking_calls import (
    calculate_call_schedule,
    BookingCallsManager,
    CALL_STATUSES
)


class TestBookingCallsScheduling(unittest.TestCase):
    def test_rule_1_2_3_4_seven_night_stay(self):
        # Arrival: 10/09/2026, Departure: 17/09/2026 (7 nights)
        # Rule 1 (A+1): 11
        # Rule 2 (D-1): 16
        # Rule 3 (every other day): 13, 15
        # Candidates: [11, 13, 15, 16]
        # Rule 4: 15 and 16 are consecutive! Discard earlier (15), keep later (16).
        # Expected: [11, 13, 16]
        arr = date(2026, 9, 10)
        dep = date(2026, 9, 17)
        schedule = calculate_call_schedule(arr, dep)
        expected = [date(2026, 9, 11), date(2026, 9, 13), date(2026, 9, 16)]
        self.assertEqual(schedule, expected)

    def test_eight_night_stay_no_consecutive(self):
        # Arrival: 10/09/2026, Departure: 18/09/2026 (8 nights)
        # Rule 1: 11
        # Rule 2: 17
        # Rule 3: 13, 15
        # Candidates: [11, 13, 15, 17]
        # Rule 4: No consecutive days
        # Expected: [11, 13, 15, 17]
        arr = date(2026, 9, 10)
        dep = date(2026, 9, 18)
        schedule = calculate_call_schedule(arr, dep)
        expected = [date(2026, 9, 11), date(2026, 9, 13), date(2026, 9, 15), date(2026, 9, 17)]
        self.assertEqual(schedule, expected)

    def test_three_night_stay(self):
        # Arrival: 10/09/2026, Departure: 13/09/2026 (3 nights)
        # Rule 1: 11
        # Rule 2: 12
        # Candidates: [11, 12] (consecutive!)
        # Rule 4: Discard 11, keep 12.
        # Expected: [12]
        arr = date(2026, 9, 10)
        dep = date(2026, 9, 13)
        schedule = calculate_call_schedule(arr, dep)
        expected = [date(2026, 9, 12)]
        self.assertEqual(schedule, expected)

    def test_two_night_stay(self):
        # Arrival: 10/09/2026, Departure: 12/09/2026 (2 nights)
        # A+1 == D-1 == 11
        # Expected: [11]
        arr = date(2026, 9, 10)
        dep = date(2026, 9, 12)
        schedule = calculate_call_schedule(arr, dep)
        expected = [date(2026, 9, 11)]
        self.assertEqual(schedule, expected)

    def test_five_night_stay(self):
        # Arrival: 10/09/2026, Departure: 15/09/2026 (5 nights)
        # Rule 1: 11
        # Rule 2: 14
        # Rule 3: 13
        # Candidates: [11, 13, 14] -> 13 and 14 consecutive! Discard 13, keep 14.
        # Expected: [11, 14]
        arr = date(2026, 9, 10)
        dep = date(2026, 9, 15)
        schedule = calculate_call_schedule(arr, dep)
        expected = [date(2026, 9, 11), date(2026, 9, 14)]
        self.assertEqual(schedule, expected)

    def test_four_night_stay(self):
        # Arrival: 10/09/2026, Departure: 14/09/2026 (4 nights)
        # Rule 1: 11
        # Rule 2: 13
        # Candidates: [11, 13]
        # Expected: [11, 13]
        arr = date(2026, 9, 10)
        dep = date(2026, 9, 14)
        schedule = calculate_call_schedule(arr, dep)
        expected = [date(2026, 9, 11), date(2026, 9, 13)]
        self.assertEqual(schedule, expected)

    def test_standard_six_statuses(self):
        self.assertEqual(len(CALL_STATUSES), 6)
        self.assertIn("Green", CALL_STATUSES)
        self.assertIn("Red", CALL_STATUSES)
        self.assertIn("Yellow", CALL_STATUSES)
        self.assertIn("N/A (NO ANSWER)", CALL_STATUSES)
        self.assertIn("N/E (NO ENGLISH)", CALL_STATUSES)
        self.assertIn("N/W (LINE NOT WORKING)", CALL_STATUSES)


class TestBookingCallsManager(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp(prefix="test_booking_calls_")
        self.xlsx_path = os.path.join(self.test_dir, "BOOKING CALLS.xlsx")
        self.state_path = os.path.join(self.test_dir, "booking_calls_state.json")

        # Create dummy workbook with 3 sheets
        wb = openpyxl.Workbook()
        ws_arr = wb.active
        ws_arr.title = "ARRIVALS"
        ws_arr.cell(1, 1).value = "DATE"

        ws_f1 = wb.create_sheet("FOLLOW UP 1")
        ws_f1.cell(1, 1).value = "DATE"

        ws_f2 = wb.create_sheet("FOLLOW UP")
        ws_f2.cell(1, 1).value = "DATE"

        wb.save(self.xlsx_path)

        self.today_json_dir = os.path.join(self.test_dir, "booking calls for today")
        self.today_json_path = os.path.join(self.today_json_dir, "booking_calls_today.json")

        self.manager = BookingCallsManager(
            xlsx_path=self.xlsx_path,
            state_path=self.state_path,
            today_json_path=self.today_json_path
        )

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_is_booking_com_filtering(self):
        from booking_calls import is_booking_com
        # Matches Booking.com in various fields
        self.assertTrue(is_booking_com({"Χρεώστης": "BOOKING.COM"}))
        self.assertTrue(is_booking_com({"Χρεώστης": "Booking.com B.V."}))
        self.assertTrue(is_booking_com({"agency": "booking.com"}))
        self.assertTrue(is_booking_com({"Τιμοκατάλογος": "BOOKING.COM BAR 2026"}))

        # Strictly rejects all other agencies
        self.assertFalse(is_booking_com({"Χρεώστης": "TUI DEUTSCHLAND"}))
        self.assertFalse(is_booking_com({"Χρεώστης": "EXPEDIA"}))
        self.assertFalse(is_booking_com({"Χρεώστης": "DERTOUR GMBH"}))
        self.assertFalse(is_booking_com({"Χρεώστης": "RAINBOW TOURS"}))
        self.assertFalse(is_booking_com({}))

    def test_state_persistence(self):
        self.manager.save_room_call(
            room="1101",
            status="Green",
            notes="Guest loves the sea view",
            call_date=date(2026, 9, 14)
        )

        state = self.manager.load_calls_state()
        self.assertIn("2026-09-14", state)
        self.assertIn("1101", state["2026-09-14"])
        self.assertEqual(state["2026-09-14"]["1101"]["status"], "Green")
        self.assertEqual(state["2026-09-14"]["1101"]["notes"], "Guest loves the sea view")

    def test_today_calls_json_strictly_in_designated_directory(self):
        sample_calls = [
            {
                "room": "1205",
                "property": "SANDY BEACH",
                "booking_id": "99123",
                "guest_name": "Smith John",
                "agency": "BOOKING.COM",
                "arrival": "10/09/2026",
                "departure": "17/09/2026",
                "status": "Green",
                "notes": "First call OK",
                "date": "2026-09-14"
            }
        ]
        self.manager.save_today_calls_json(sample_calls)

        self.assertTrue(os.path.exists(self.today_json_path))
        loaded = self.manager.load_today_calls_json()
        self.assertEqual(len(loaded), 1)
        self.assertEqual(loaded[0]["room"], "1205")
        self.assertEqual(loaded[0]["agency"], "BOOKING.COM")
        self.assertEqual(loaded[0]["status"], "Green")

    def test_populate_strictly_in_followup_sheet_never_sheet_1(self):
        # Mock master state with a Booking.com booking and a non-Booking.com booking
        mock_master = {
            "10001": {
                "Δωμάτιο": "1101",
                "Άφιξη": "10/09/2026",
                "Αναχώρηση": "17/09/2026",
                "agency": "BOOKING.COM",
                "Πελάτες": ["Booking Guest"]
            },
            "20002": {
                "Δωμάτιο": "1202",
                "Άφιξη": "10/09/2026",
                "Αναχώρηση": "17/09/2026",
                "agency": "TUI",
                "Πελάτες": ["TUI Guest"]
            }
        }
        self.manager.data_manager.load_master_state = lambda: mock_master

        # Run follow-up synchronization
        schedule = self.manager.sync_followup_schedule_to_excel()
        self.assertGreater(len(schedule), 0)

        wb = openpyxl.load_workbook(self.xlsx_path)

        # 1. Sheet 1 (first sheet) MUST NOT have booking calls data populated
        first_sheet = wb.worksheets[0]
        # Any row 3 onwards in first sheet should remain None
        for r in range(3, max(first_sheet.max_row + 1, 10)):
            for c in range(2, max(first_sheet.max_column + 1, 5)):
                self.assertIsNone(first_sheet.cell(r, c).value)

        # 2. Follow-Up Sheet MUST have Booking.com room 1101 populated
        followup_sheet = self.manager._get_followup_sheet(wb)
        self.assertIsNotNone(followup_sheet)
        self.assertNotEqual(followup_sheet.title, first_sheet.title)

        found_1101 = False
        found_1202 = False
        for r in range(3, followup_sheet.max_row + 1):
            for c in range(1, followup_sheet.max_column + 1):
                val = str(followup_sheet.cell(r, c).value or "")
                if val == "1101":
                    found_1101 = True
                if val == "1202":
                    found_1202 = True

        self.assertTrue(found_1101, "Booking.com room 1101 must be in Follow-Up sheet")
        self.assertFalse(found_1202, "Non-Booking.com room 1202 must NOT be in Follow-Up sheet")

    def test_feedback_to_do_formatting(self):
        comment = "Guest requested extra pillows"
        room = "1402"
        expected_format = f"Feedback on exclusivi: {comment} - Room {room}"
        self.assertEqual(expected_format, "Feedback on exclusivi: Guest requested extra pillows - Room 1402")


if __name__ == "__main__":
    unittest.main()

