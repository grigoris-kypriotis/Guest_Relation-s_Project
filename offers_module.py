import os
import glob
import csv
import re
import shutil
import traceback
from datetime import datetime
import win32com.client
import pythoncom

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

def _load_env_file():
    env_path = os.path.join(BASE_DIR, ".env")
    if os.path.exists(env_path):
        try:
            with open(env_path, "r", encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if line and not line.startswith("#") and "=" in line:
                        k, v = line.split("=", 1)
                        k = k.strip()
                        v = v.strip().strip('"').strip("'")
                        if k not in os.environ:
                            os.environ[k] = v
        except Exception:
            pass

_load_env_file()

from data_manager import resolve_template_path, DATABASE_DIR, OUTPUT_DIR, TODAYS_LIST_DIR

ARRIVALS_FOLDER = os.path.abspath(DATABASE_DIR)
FINAL_FOLDER = os.path.abspath(os.path.join(OUTPUT_DIR, "OFFERS"))
TODAY_LIST_FOLDER = os.path.abspath(TODAYS_LIST_DIR)

TEMPLATE_PATH = resolve_template_path("offer_list") or os.path.abspath(os.path.join(BASE_DIR, "templates", "offer list template", "OFFER LIST TEMPLATE.docx"))

CHECK_MEMO_TEMPLATE_PATH = resolve_template_path("check_memo") or os.path.abspath(os.path.join(BASE_DIR, "templates", "check memo template", "CHECK MEMO.docx"))
CAKE_TEMPLATE_PATH = CHECK_MEMO_TEMPLATE_PATH
CAKE_MEMO_TEMPLATE_PATH = CHECK_MEMO_TEMPLATE_PATH

os.makedirs(FINAL_FOLDER, exist_ok=True)
os.makedirs(TODAY_LIST_FOLDER, exist_ok=True)

def log_task(task_text, status):
    os.makedirs(TODAY_LIST_FOLDER, exist_ok=True)
    now = datetime.now()
    filepath = os.path.abspath(os.path.join(TODAY_LIST_FOLDER, f"Tasks_{now.strftime('%Y-%m-%d')}.txt"))
    with open(filepath, 'a', encoding='utf-8') as f:
        f.write(f"[{now.strftime('%H:%M:%S')}] {status}: {task_text}\n")

class OutlookMailEvents:
    def OnSend(self, Cancel):
        if hasattr(self, 'callback') and hasattr(self, 'task_id'):
            self.callback(self.task_id, "SENT")
    def OnClose(self, Cancel):
        if hasattr(self, 'callback') and hasattr(self, 'task_id'):
            self.callback(self.task_id, "CLOSED")

def identify_digit_type(file_path):
    with open(file_path, mode='r', encoding='utf-8-sig', errors='ignore') as f:
        reader = csv.reader(f, delimiter=';')
        for i, row in enumerate(reader):
            if i >= 50: break
            if not row: continue
            col_a = row[0].strip(' "\'')
            if len(col_a) == 3 and col_a.isdigit(): return 3
            if len(col_a) == 4 and col_a.isdigit(): return 4
    return 0

def extract_excel_data(csv_path, is_villas):
    csv_path = os.path.abspath(csv_path)
    extracted_data = []
    minibar_data = []
    
    with open(csv_path, mode='r', encoding='utf-8-sig', errors='ignore') as f:
        rows = list(csv.reader(f, delimiter=';'))
        
    row_count = len(rows)
    for r in range(row_count):
        row = rows[r]
        next_row = rows[r + 1] if (r + 1) < row_count else []
        
        col_a = row[0].strip() if len(row) > 0 else ""
        col_e_text = row[4].strip() if len(row) > 4 else ""
        col_i = row[8].strip() if len(row) > 8 else ""
        desc = next_row[0].strip() if len(next_row) > 0 else ""
        dep_date = ""
        
        if col_e_text:
            date_match = re.search(r"(\d{1,2})[\/\-\.](\d{1,2})", col_e_text)
            if date_match:
                dep_date = f"{int(date_match.group(1)):02d}/{int(date_match.group(2)):02d}"
            else:
                dep_date = col_e_text.strip()
            
        is_booking = bool(re.search(r"(?i)BOOKING\.COM", col_i))
        is_room = bool(re.match(r"^\d{3,4}$", col_a))
        
        if is_booking:
            room_no = col_a.split()[0] if (not is_room and col_a) else col_a
            if not room_no: continue
            pax = sum(int(re.sub(r'\D', '', row[c])) for c in range(9, min(15, len(row))) if re.search(r'\d', row[c]))
            extracted_data.append({"RoomNo": room_no, "DepDate": dep_date, "Pax": pax, "Order": "ST"})
            
        elif is_room:
            if not col_a: continue
            order_str = None
            if re.search(r"(?i)Anniversary|Birthday|Honeymoon|Brthd|VIP", desc): order_str = "HB"
            elif re.search(r"(?i)Fruit", desc): order_str = "ST"
                
            if order_str:
                pax = sum(int(re.sub(r'\D', '', row[c])) for c in range(9, min(15, len(row))) if re.search(r'\d', row[c]))
                extracted_data.append({"RoomNo": col_a, "DepDate": dep_date, "Pax": pax, "Order": order_str})
                
        if is_villas and is_room and len(col_a) == 3:
            if int(col_a) >= 600: minibar_data.append(col_a)
            elif re.search(r"(?i)mini bar|refill", desc): minibar_data.append(f"{col_a} REFILL")
            
    return extracted_data, minibar_data

def generate_word_document(data, minibar_data, today_docx, save_path, year_str):
    save_path = os.path.abspath(save_path)
    
    # Primary generator: python-docx (fast, fully portable, no MS Office installation required)
    try:
        import docx
        doc = docx.Document(os.path.abspath(TEMPLATE_PATH))
        
        # Replace delivery date placeholder
        for p in doc.paragraphs:
            if "Delivery Date: / /" in p.text:
                p.text = p.text.replace("Delivery Date: / /", f"Delivery Date: {today_docx}/{year_str}")
                
        if doc.tables:
            table = doc.tables[0]
            col_indices = {}
            for idx, c in enumerate(table.rows[0].cells):
                header = re.sub(r'[\r\a\n\t]', '', c.text).strip()
                if re.search(r"(?i)Room No", header): col_indices["Room No."] = idx
                if re.search(r"(?i)Dep Date", header): col_indices["Dep Date"] = idx
                if re.search(r"(?i)Pax", header): col_indices["Pax"] = idx
                if re.search(r"(?i)Delivery Date", header): col_indices["Delivery Date"] = idx
                if re.search(r"(?i)Order", header): col_indices["Order"] = idx
                if re.search(r"(?i)Mini bar", header): col_indices["Mini bar"] = idx
                
            current_row = 1
            for item in data:
                if current_row < len(table.rows):
                    row = table.rows[current_row]
                else:
                    row = table.add_row()
                if "Room No." in col_indices: row.cells[col_indices["Room No."]].text = str(item["RoomNo"])
                if "Dep Date" in col_indices: row.cells[col_indices["Dep Date"]].text = str(item["DepDate"])
                if "Pax" in col_indices: row.cells[col_indices["Pax"]].text = str(item["Pax"])
                if "Delivery Date" in col_indices: row.cells[col_indices["Delivery Date"]].text = today_docx
                if "Order" in col_indices: row.cells[col_indices["Order"]].text = str(item["Order"])
                current_row += 1
                
            if "Mini bar" in col_indices:
                mb_row = 1
                for room in minibar_data:
                    if mb_row < len(table.rows):
                        row = table.rows[mb_row]
                    else:
                        row = table.add_row()
                    row.cells[col_indices["Mini bar"]].text = str(room)
                    mb_row += 1
                    
        doc.save(save_path)
        return
    except Exception as docx_err:
        pass

    # Fallback to win32com Word.Application if docx fails or is not available
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    doc = None
    try:
        doc = word.Documents.Open(os.path.abspath(TEMPLATE_PATH))
        word.Selection.HomeKey(Unit=6)
        word.Selection.Find.Execute(FindText="Delivery Date: / /", ReplaceWith=f"Delivery Date: {today_docx}/{year_str}", Replace=2)
        
        table = doc.Tables(1)
        col_indices = {}
        for c in range(1, table.Columns.Count + 1):
            header = re.sub(r'[\r\a\n\t]', '', table.Cell(1, c).Range.Text).strip()
            if re.search(r"(?i)Room No", header): col_indices["Room No."] = c
            if re.search(r"(?i)Dep Date", header): col_indices["Dep Date"] = c
            if re.search(r"(?i)Pax", header): col_indices["Pax"] = c
            if re.search(r"(?i)Delivery Date", header): col_indices["Delivery Date"] = c
            if re.search(r"(?i)Order", header): col_indices["Order"] = c
            if re.search(r"(?i)Mini bar", header): col_indices["Mini bar"] = c
            
        current_row = 2
        for item in data:
            if current_row > table.Rows.Count: table.Rows.Add()
            if "Room No." in col_indices: table.Cell(current_row, col_indices["Room No."]).Range.Text = str(item["RoomNo"])
            if "Dep Date" in col_indices: table.Cell(current_row, col_indices["Dep Date"]).Range.Text = str(item["DepDate"])
            if "Pax" in col_indices: table.Cell(current_row, col_indices["Pax"]).Range.Text = str(item["Pax"])
            if "Delivery Date" in col_indices: table.Cell(current_row, col_indices["Delivery Date"]).Range.Text = today_docx
            if "Order" in col_indices: table.Cell(current_row, col_indices["Order"]).Range.Text = str(item["Order"])
            current_row += 1
            
        if "Mini bar" in col_indices:
            mb_row = 2
            for room in minibar_data:
                if mb_row > table.Rows.Count: table.Rows.Add()
                table.Cell(mb_row, col_indices["Mini bar"]).Range.Text = str(room)
                mb_row += 1
                
        doc.SaveAs2(save_path)
    finally:
        if doc:
            try: doc.Close(False)
            except: pass
        try: word.Quit()
        except: pass

def draft_email_payload(file_path, is_updated, task_id, callback):
    file_path = os.path.abspath(file_path)
    outlook = win32com.client.Dispatch("Outlook.Application")
    mail = win32com.client.DispatchWithEvents(outlook.CreateItem(0), OutlookMailEvents)
    mail.task_id = task_id
    mail.callback = callback
    
    # Configurable email recipients with generic placeholders (configure in environment or local .env)
    default_to = (
        "Executive Chef <chef@example.com>; "
        "headchef@example.com; "
        "F&B Manager <fb.manager@example.com>; "
        "assistfb@example.com"
    )
    default_cc = (
        "Operation Manager <operations@example.com>; "
        "Rooms Division Manager <rooms@example.com>; "
        "Front Office Manager <frontoffice@example.com>; "
        "Guest Relations <guestrelations@example.com>"
    )
    mail.To = os.getenv("OFFERS_MAIL_TO", default_to)
    mail.CC = os.getenv("OFFERS_MAIL_CC", default_cc)
    mail.Subject = os.path.basename(file_path).replace(".docx", "")
    mail.Attachments.Add(file_path)
    mail.Display()
    
    now = datetime.now()
    date_dd = now.strftime("%d")
    date_mm = now.strftime("%m")
    
    if is_updated:
        body = "Dear all,<br>Kindly find attached the updated version of today’s offer list.<br><br><br>For any further information don’t hesitate to contact the Guest Relations Team."
    else:
        body = f"Dear all,<br>I have attached todays offers list {date_dd}/{date_mm}.<br><br>For any further information don’t hesitate to speak with the Guest Relations Team."
        
    html_content = f"<div style='font-family: Calibri, sans-serif; font-size: 11pt;'>{body}<br><br></div>"
    mail.HTMLBody = html_content + mail.HTMLBody
    return mail

def duplicate_for_update(file_path):
    file_path = os.path.abspath(file_path)
    dir_name = os.path.dirname(file_path)
    base_name = os.path.basename(file_path).replace(".docx", "")
    base_name = re.sub(r" UPDATED(?: \(\d+\))?", "", base_name)
    
    new_path = os.path.abspath(os.path.join(dir_name, base_name + " UPDATED.docx"))
    counter = 2
    while os.path.exists(new_path):
        new_path = os.path.abspath(os.path.join(dir_name, base_name + f" UPDATED ({counter}).docx"))
        counter += 1
        
    shutil.copy2(file_path, new_path)
    return new_path

def get_todays_offer_list():
    now = datetime.now()
    target_folder = os.path.abspath(os.path.join(FINAL_FOLDER, f"GR OFFERS {now.month}.{now.year}"))
    base_name = f"OFFER LIST ({now.strftime('%Y-%m-%d')})"
    search_pattern = os.path.abspath(os.path.join(target_folder, f"{base_name}*.docx"))
    files = glob.glob(search_pattern)
    if not files: return None
    return os.path.abspath(max(files, key=os.path.getmtime))

def execute_offers_pipeline(selected_csvs=None):
    try:
        if not os.path.exists(TEMPLATE_PATH):
            return False, f"Template missing at {TEMPLATE_PATH}", None
            
        if selected_csvs and len(selected_csvs) == 2:
            csv_files = [os.path.abspath(f) for f in selected_csvs]
        else:
            csv_files = [os.path.abspath(f) for f in glob.glob(os.path.join(ARRIVALS_FOLDER, "*.csv"))]
            if len(csv_files) != 2:
                beach_sub = glob.glob(os.path.join(ARRIVALS_FOLDER, "SANDY BEACH", "*.csv"))
                villas_sub = glob.glob(os.path.join(ARRIVALS_FOLDER, "SANDY VILLAS", "*.csv"))
                if beach_sub and villas_sub:
                    latest_beach = max(beach_sub, key=os.path.getmtime)
                    latest_villas = max(villas_sub, key=os.path.getmtime)
                    csv_files = [os.path.abspath(latest_beach), os.path.abspath(latest_villas)]
                else:
                    sub_csvs = [os.path.abspath(f) for f in glob.glob(os.path.join(ARRIVALS_FOLDER, "**", "*.csv"), recursive=True)]
                    if len(sub_csvs) == 2:
                        csv_files = sub_csvs
            
        if len(csv_files) != 2: return False, "MISSING_CSVS", None
            
        t1 = identify_digit_type(csv_files[0])
        t2 = identify_digit_type(csv_files[1])
        
        if (t1 == 3 and t2 == 3) or (t1 == 4 and t2 == 4) or t1 == 0 or t2 == 0:
            return False, "CSV validation failed.", None
            
        villas_csv = csv_files[0] if t1 == 3 else csv_files[1]
        beach_csv = csv_files[0] if t1 == 4 else csv_files[1]
        
        beach_data, _ = extract_excel_data(beach_csv, is_villas=False)
        villas_data, minibar_data = extract_excel_data(villas_csv, is_villas=True)
        
        villas_dir = os.path.abspath(os.path.join(ARRIVALS_FOLDER, "SANDY VILLAS"))
        beach_dir = os.path.abspath(os.path.join(ARRIVALS_FOLDER, "SANDY BEACH"))
        os.makedirs(villas_dir, exist_ok=True)
        os.makedirs(beach_dir, exist_ok=True)
        
        dest_villas = os.path.abspath(os.path.join(villas_dir, os.path.basename(villas_csv)))
        if os.path.abspath(villas_csv) != dest_villas:
            shutil.move(villas_csv, dest_villas)
            villas_csv = dest_villas
            
        dest_beach = os.path.abspath(os.path.join(beach_dir, os.path.basename(beach_csv)))
        if os.path.abspath(beach_csv) != dest_beach:
            shutil.move(beach_csv, dest_beach)
            beach_csv = dest_beach
        
        all_data = sorted(beach_data + villas_data, key=lambda x: (not str(x["RoomNo"]).isdigit(), int(x["RoomNo"]) if str(x["RoomNo"]).isdigit() else str(x["RoomNo"])))
        
        now = datetime.now()
        today_docx = now.strftime("%d/%m")
        year_str = now.strftime("%Y")
        target_folder = os.path.abspath(os.path.join(FINAL_FOLDER, f"GR OFFERS {now.month}.{year_str}"))
        os.makedirs(target_folder, exist_ok=True)
        
        base_name = f"OFFER LIST ({now.strftime('%Y-%m-%d')})"
        final_path = os.path.abspath(os.path.join(target_folder, base_name + ".docx"))
        
        if os.path.exists(final_path):
            final_path = os.path.abspath(os.path.join(target_folder, base_name + " UPDATED.docx"))
            counter = 2
            while os.path.exists(final_path):
                final_path = os.path.abspath(os.path.join(target_folder, base_name + f" UPDATED ({counter}).docx"))
                counter += 1
                
        generate_word_document(all_data, minibar_data, today_docx, final_path, year_str)
        return True, "Success: Pipeline complete.", final_path
        
    except Exception as e:
        error_msg = f"Pipeline Error: {str(e)}\n\nTraceback:\n{traceback.format_exc()}"
        return False, error_msg, None