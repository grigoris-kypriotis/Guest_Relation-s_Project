"""
Booking Calls CRM Module & Stay-Based Follow-Up Scheduling Engine
================================================================
Strictly processes ONLY reservations where the travel agent/booker is BOOKING.COM.
Completely ignores all other agencies.

Features:
1. Stay-based Calling Schedule Algorithm (Rule 1, 2, 3, 4) strictly for Booking.com
2. Synchronization of BOOKING CALLS.xlsx:
   - Sheet 1 (ARRIVALS): Populated exclusively with today's Booking.com arrivals
   - Sheet 2 (FOLLOW UP): Populated with Booking.com rooms scheduled per date
3. Direct Cell Manipulation in Excel:
   - Status Color Highlighting: Green (#D4EDDA), Yellow (#FFF3CD), Red (#F8D7DA)
   - Status String Insertion: 'room N/A', 'room N/E', 'room N/W' directly in the same cell
4. Centralized Directory & State Management:
   - Today's calls JSON saved strictly inside 'DATABASE/BOOKING CALLS FOR TODAY/booking_calls_for_today.json'
5. Interactive PyQt6 CRM interface with editable room card containers and 6 standardized statuses.
"""

import os
import re
import json
from datetime import datetime, date, timedelta
from typing import Dict, List, Tuple, Optional, Any

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QComboBox, QTextEdit, QLineEdit, QScrollArea, QFrame,
    QGridLayout, QMessageBox, QSizePolicy, QToolButton
)
from PyQt6.QtCore import Qt, pyqtSignal, QTimer
from PyQt6.QtGui import QFont, QColor

from data_manager import (
    InHouseDataManager, BASE_DIR, DATABASE_DIR, OUTPUT_DIR,
    BOOKING_CALLS_TODAY_DIR, BOOKING_CALLS_TODAY_JSON,
    save_and_archive_json, resolve_template_path
)

BOOKING_CALLS_DIR = os.path.join(BASE_DIR, "BOOKING CALLS")
BOOKING_CALLS_XLSX = os.path.join(BOOKING_CALLS_DIR, "BOOKING CALLS.xlsx")
TEMPLATE_BOOKING_CALLS_XLSX = resolve_template_path("booking_calls") or os.path.join(BASE_DIR, "TEMPLATES", "BOOKING CALLS TEMPLATE", "BOOKING CALLS.xlsx")
BOOKING_CALLS_STATE_PATH = os.path.join(BOOKING_CALLS_TODAY_DIR, "booking_calls_state.json")

os.makedirs(BOOKING_CALLS_DIR, exist_ok=True)
os.makedirs(BOOKING_CALLS_TODAY_DIR, exist_ok=True)

# Exactly 6 Standardized Status Options
CALL_STATUSES = [
    "Green",
    "Red",
    "Yellow",
    "N/A (NO ANSWER)",
    "N/E (NO ENGLISH)",
    "N/W (LINE NOT WORKING)"
]

STATUS_COLORS = {
    "Green": {"bg": "#D4EDDA", "border": "#28A745", "text": "#155724", "excel_hex": "D4EDDA"},
    "Red": {"bg": "#F8D7DA", "border": "#DC3545", "text": "#721C24", "excel_hex": "F8D7DA"},
    "Yellow": {"bg": "#FFF3CD", "border": "#FFC107", "text": "#856404", "excel_hex": "FFF3CD"},
    "N/A (NO ANSWER)": {"bg": "#E2E3E5", "border": "#6C757D", "text": "#383D41", "excel_hex": "E2E3E5", "code": "N/A"},
    "N/E (NO ENGLISH)": {"bg": "#D1ECF1", "border": "#17A2B8", "text": "#0C5460", "excel_hex": "D1ECF1", "code": "N/E"},
    "N/W (LINE NOT WORKING)": {"bg": "#F5C6CB", "border": "#C82333", "text": "#491217", "excel_hex": "F5C6CB", "code": "N/W"}
}


# =============================================================================
# Agency Filter: STRICTLY BOOKING.COM ONLY
# =============================================================================

def is_booking_com(reservation_data: Dict[str, Any]) -> bool:
    """
    Returns True if and only if the reservation's travel agent / booker is BOOKING.COM.
    Completely filters out all other agencies (TUI, Expedia, Rainbow, etc.).
    """
    if not isinstance(reservation_data, dict):
        return False
    debtor = str(reservation_data.get("Χρεώστης", "")).strip()
    rate_plan = str(reservation_data.get("Τιμοκατάλογος", "")).strip()
    agency = str(reservation_data.get("agency", "")).strip()

    combined = f"{debtor} {rate_plan} {agency}".upper()
    return bool(re.search(r"(?i)BOOKING\.COM", combined))


# =============================================================================
# Scheduling Algorithm (Rule 1, 2, 3, 4)
# =============================================================================

def calculate_call_schedule(arrival_date: date, departure_date: date) -> List[date]:
    """
    Calculates the calling schedule for a guest stay based on the 4 mandatory rules:
      Rule 1: Call 1 day after arrival (A + 1).
      Rule 2: Call 1 day prior to departure (D - 1).
      Rule 3: Call every other day in between (A + 3, A + 5, ... up to D - 2).
      Rule 4: Never schedule calls for two consecutive days. If the algorithm results
              in two consecutive days, strictly keep the later date and discard the earlier one.
    """
    if departure_date <= arrival_date:
        return []

    day_after_arr = arrival_date + timedelta(days=1)
    day_before_dep = departure_date - timedelta(days=1)

    if day_after_arr > day_before_dep:
        return []

    if day_after_arr == day_before_dep:
        return [day_after_arr]

    candidate_dates = set()

    # Rule 1: Call 1 day after arrival
    candidate_dates.add(day_after_arr)

    # Rule 2: Call 1 day prior to departure
    candidate_dates.add(day_before_dep)

    # Rule 3: Call every other day in between
    curr = day_after_arr + timedelta(days=2)
    while curr < day_before_dep:
        candidate_dates.add(curr)
        curr += timedelta(days=2)

    sorted_dates = sorted(list(candidate_dates))

    # Rule 4: Never schedule calls for two consecutive days.
    # If the algorithm results in two consecutive days, strictly keep the later date
    # and discard the earlier one.
    resolved = []
    i = 0
    while i < len(sorted_dates):
        if i + 1 < len(sorted_dates) and (sorted_dates[i + 1] - sorted_dates[i]).days == 1:
            i += 1  # Skip the earlier date, keep the later date
        else:
            resolved.append(sorted_dates[i])
            i += 1

    return resolved


def parse_date_str(date_str: str) -> Optional[date]:
    """Helper to parse dates in various formats (d/m/Y, Y-m-d, etc.)."""
    if not date_str:
        return None
    cleaned = date_str.strip().split()[0]
    for fmt in ["%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"]:
        try:
            return datetime.strptime(cleaned, fmt).date()
        except ValueError:
            pass
    return None


# =============================================================================
# Booking Calls Manager Engine & Excel Synchronization
# =============================================================================

class BookingCallsManager:
    """
    Core business logic for the Booking Calls CRM module.
    Handles scheduling, Excel updates for Sheet 1 (Arrivals) and Sheet 2 (Follow-Up),
    direct cell manipulation (color highlight & status string insertion), and JSON persistence.
    """

    def __init__(
        self,
        data_manager: Optional[InHouseDataManager] = None,
        xlsx_path: str = BOOKING_CALLS_XLSX,
        state_path: str = BOOKING_CALLS_STATE_PATH,
        today_json_path: str = BOOKING_CALLS_TODAY_JSON
    ):
        self.data_manager = data_manager or InHouseDataManager()
        self.xlsx_path = os.path.abspath(xlsx_path)
        self.state_path = os.path.abspath(state_path)
        self.today_json_path = os.path.abspath(today_json_path)

    # -------------------------------------------------------------------------
    # State Persistence
    # -------------------------------------------------------------------------
    def load_calls_state(self) -> Dict[str, Any]:
        """Loads master booking calls state (call history, notes, status)."""
        if os.path.exists(self.state_path):
            try:
                with open(self.state_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    if isinstance(data, dict):
                        return data
            except Exception as e:
                print(f"[BookingCallsManager] Error reading calls state: {e}")
        return {}

    def save_calls_state(self, state: Dict[str, Any]):
        """Persists master calls state."""
        try:
            target_fname = os.path.basename(self.state_path)
            if self.state_path == BOOKING_CALLS_STATE_PATH:
                save_and_archive_json(state, target_fname, subfolder="BOOKING CALLS FOR TODAY")
            else:
                os.makedirs(os.path.dirname(self.state_path), exist_ok=True)
                temp_path = self.state_path + ".tmp"
                with open(temp_path, "w", encoding="utf-8") as f:
                    json.dump(state, f, ensure_ascii=False, indent=4)
                if os.path.exists(self.state_path):
                    os.replace(temp_path, self.state_path)
                else:
                    os.rename(temp_path, self.state_path)
        except Exception as e:
            print(f"[BookingCallsManager] Error saving calls state: {e}")

    # -------------------------------------------------------------------------
    # Today's Specific Calls JSON
    # -------------------------------------------------------------------------
    def load_today_calls_json(self) -> List[Dict[str, Any]]:
        """Reads today's scheduled calls strictly from BOOKING CALLS FOR TODAY/booking_calls_for_today.json."""
        if os.path.exists(self.today_json_path):
            try:
                with open(self.today_json_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                    if isinstance(data, list):
                        return data
                    if isinstance(data, dict) and "calls" in data:
                        return data["calls"]
            except Exception as e:
                print(f"[BookingCallsManager] Error reading today JSON: {e}")

        return self.generate_today_booking_calls_json()

    def save_today_calls_json(self, calls_list: List[Dict[str, Any]]):
        """Strictly saves today's booking calls into DATABASE/BOOKING CALLS FOR TODAY/booking_calls_for_today.json."""
        try:
            target_fname = os.path.basename(self.today_json_path)
            if self.today_json_path == BOOKING_CALLS_TODAY_JSON:
                save_and_archive_json(calls_list, target_fname, subfolder="BOOKING CALLS FOR TODAY")
            else:
                os.makedirs(os.path.dirname(self.today_json_path), exist_ok=True)
                temp_path = self.today_json_path + ".tmp"
                with open(temp_path, "w", encoding="utf-8") as f:
                    json.dump(calls_list, f, ensure_ascii=False, indent=4)
                if os.path.exists(self.today_json_path):
                    os.replace(temp_path, self.today_json_path)
                else:
                    os.rename(temp_path, self.today_json_path)
        except Exception as e:
            print(f"[BookingCallsManager] Error saving today calls JSON: {e}")

    # -------------------------------------------------------------------------
    # Excel Workbook Helpers
    # -------------------------------------------------------------------------
    def _get_or_create_workbook(self) -> openpyxl.Workbook:
        """Loads BOOKING CALLS.xlsx or copies it from TEMPLATES or creates it with default sheets."""
        if os.path.exists(self.xlsx_path):
            return openpyxl.load_workbook(self.xlsx_path)

        resolved_tpl = resolve_template_path("booking_calls") or TEMPLATE_BOOKING_CALLS_XLSX
        if resolved_tpl and os.path.exists(resolved_tpl):
            try:
                import shutil
                shutil.copy2(resolved_tpl, self.xlsx_path)
                return openpyxl.load_workbook(self.xlsx_path)
            except Exception as e:
                print(f"[BookingCallsManager] Error copying template from {resolved_tpl}: {e}")

        wb = openpyxl.Workbook()
        ws_arr = wb.active
        ws_arr.title = "ARRIVALS"
        ws_arr.cell(1, 1).value = "DATE"

        ws_fu = wb.create_sheet("FOLLOW UP")
        ws_fu.cell(1, 1).value = "DATE"
        try:
            wb.save(self.xlsx_path)
        except PermissionError:
            raise PermissionError(
                f"Cannot write to '{os.path.basename(self.xlsx_path)}'. "
                "Please close Microsoft Excel or any program accessing this file, then try again."
            )
        return wb

    def _get_sheet(self, wb: openpyxl.Workbook, sheet_type: str):
        """Locates the designated sheet ('ARRIVALS' or 'FOLLOW UP')."""
        target = sheet_type.upper().strip()
        for name in wb.sheetnames:
            if target == name.strip().upper():
                return wb[name]
        for name in wb.sheetnames:
            if target in name.strip().upper():
                return wb[name]
        if target == "ARRIVALS":
            return wb.worksheets[0]
        else:
            if len(wb.sheetnames) > 1:
                return wb.worksheets[1]
            return wb.create_sheet("FOLLOW UP")

    def _save_workbook(self, wb: openpyxl.Workbook):
        """Safely saves workbook, catching Excel file locks."""
        try:
            wb.save(self.xlsx_path)
        except PermissionError:
            raise PermissionError(
                f"Cannot write to '{os.path.basename(self.xlsx_path)}'. "
                "Please close Microsoft Excel or any program accessing this file, then try again."
            )

    def _find_or_create_date_column(self, ws, target_date: date) -> int:
        """Finds or appends a date column in row 1."""
        target_dt = datetime(target_date.year, target_date.month, target_date.day)
        for c in range(1, ws.max_column + 1):
            val = ws.cell(1, c).value
            if isinstance(val, (datetime, date)):
                val_date = val.date() if isinstance(val, datetime) else val
                if val_date == target_date:
                    return c
            elif isinstance(val, str):
                parsed = parse_date_str(val)
                if parsed and parsed == target_date:
                    return c

        new_col = ws.max_column + 1
        ws.cell(1, new_col).value = target_dt
        ws.cell(1, new_col).number_format = 'yyyy-mm-dd'
        return new_col

    # -------------------------------------------------------------------------
    # Sheet 1: Populate Arrivals strictly for BOOKING.COM
    # -------------------------------------------------------------------------
    def sync_today_arrivals_to_excel(self, target_date: Optional[date] = None) -> List[str]:
        """
        Populates Sheet 1 (ARRIVALS) of BOOKING CALLS.xlsx strictly with
        today's arrivals where agency is BOOKING.COM.
        """
        if target_date is None:
            target_date = date.today()

        wb = self._get_or_create_workbook()
        ws = self._get_sheet(wb, "ARRIVALS")

        # Pull today's arrivals
        arrivals_state = self.data_manager.load_arrivals_state()
        beach_arrivals = arrivals_state.get("SANDY BEACH", {})

        booking_com_rooms: List[str] = []
        for b_id, arr_data in beach_arrivals.items():
            if not isinstance(arr_data, dict):
                continue
            if not is_booking_com(arr_data):
                continue
            room = str(arr_data.get("room", "")).strip()
            if room and room not in booking_com_rooms:
                booking_com_rooms.append(room)

        # Fallback to master state arrivals matching target_date if arrivals list is empty
        if not booking_com_rooms:
            master = self.data_manager.load_master_state()
            for b_id, b_data in master.items():
                if not isinstance(b_data, dict) or not is_booking_com(b_data):
                    continue
                arr_dt = parse_date_str(b_data.get("Άφιξη", ""))
                if arr_dt == target_date:
                    room = str(b_data.get("Δωμάτιο", "")).strip()
                    if room and room not in booking_com_rooms:
                        booking_com_rooms.append(room)

        # Write to date column
        col_idx = self._find_or_create_date_column(ws, target_date)
        for r in range(3, max(ws.max_row + 1, 50)):
            ws.cell(r, col_idx).value = None

        for idx, room in enumerate(booking_com_rooms):
            row_num = idx + 3
            if not ws.cell(row_num, 1).value:
                ws.cell(row_num, 1).value = idx + 1
            val = int(room) if room.isdigit() else room
            ws.cell(row_num, col_idx).value = val

        self._save_workbook(wb)
        return booking_com_rooms

    # -------------------------------------------------------------------------
    # Sheet 2: Synchronize Follow-Up Schedule - STRICTLY BOOKING.COM ONLY
    # -------------------------------------------------------------------------
    def sync_followup_schedule_to_excel(self) -> Dict[date, List[str]]:
        """
        Calculates the calling schedule strictly for active BOOKING.COM bookings
        using Rules 1, 2, 3, and 4, and populates Sheet 2 (FOLLOW UP).
        """
        wb = self._get_or_create_workbook()
        master_state = self.data_manager.load_master_state()
        date_to_rooms: Dict[date, List[str]] = {}

        for b_id, booking in master_state.items():
            if not isinstance(booking, dict) or not is_booking_com(booking):
                continue

            room = str(booking.get("Δωμάτιο", "")).strip()
            arr_date = parse_date_str(booking.get("Άφιξη", ""))
            dep_date = parse_date_str(booking.get("Αναχώρηση", ""))

            if not room or not arr_date or not dep_date:
                continue

            scheduled_dates = calculate_call_schedule(arr_date, dep_date)
            for call_dt in scheduled_dates:
                if call_dt not in date_to_rooms:
                    date_to_rooms[call_dt] = []
                if room not in date_to_rooms[call_dt]:
                    date_to_rooms[call_dt].append(room)

        ws = self._get_sheet(wb, "FOLLOW UP")

        for sched_date, rooms in date_to_rooms.items():
            col_idx = self._find_or_create_date_column(ws, sched_date)
            for r in range(3, max(ws.max_row + 1, 50)):
                ws.cell(r, col_idx).value = None

            for idx, room in enumerate(rooms):
                row_num = idx + 3
                if not ws.cell(row_num, 1).value:
                    ws.cell(row_num, 1).value = idx + 1
                val = int(room) if room.isdigit() else room
                ws.cell(row_num, col_idx).value = val

        self._save_workbook(wb)

        # Update today's calls JSON in DATABASE/BOOKING CALLS FOR TODAY/
        self.generate_today_booking_calls_json()
        return date_to_rooms

    # -------------------------------------------------------------------------
    # Direct Excel Cell Manipulation (Color Highlighting & Status String Insertion)
    # -------------------------------------------------------------------------
    def update_excel_cell_status(self, room: str, status: str, target_date: Optional[date] = None, sheet_name: Optional[str] = None) -> bool:
        """
        Directly manipulates the cell in BOOKING CALLS.xlsx for that room number:
          - Color Highlighting:
            If Green, Yellow, or Red -> sets the cell background fill:
              Green: #D4EDDA
              Yellow: #FFF3CD
              Red: #F8D7DA
          - Status String Insertion:
            If N/A (NO ANSWER), N/E (NO ENGLISH), or N/W (LINE NOT WORKING):
            writes code string directly into the SAME CELL alongside room number:
              e.g., '1024 N/A', '2015 N/E', '1102 N/W'
        """
        if not os.path.exists(self.xlsx_path):
            return False

        if target_date is None:
            target_date = date.today()

        clean_room = str(room).strip()
        if not clean_room:
            return False

        try:
            wb = openpyxl.load_workbook(self.xlsx_path)
        except Exception as e:
            print(f"[BookingCallsManager] Could not open workbook for direct cell manipulation: {e}")
            return False

        modified = False
        status_info = STATUS_COLORS.get(status, {})

        # Determine target fill
        hex_color = status_info.get("excel_hex")
        fill_to_apply = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid") if hex_color else None

        # Determine cell value
        code = status_info.get("code")
        if code:
            new_cell_value = f"{clean_room} {code}"
        else:
            new_cell_value = int(clean_room) if clean_room.isdigit() else clean_room

        # Search across target sheet or all sheets
        target_sheets = [sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.sheetnames
        for sname in target_sheets:
            ws = wb[sname]
            # Find date column if possible
            target_cols = []
            for c in range(1, ws.max_column + 1):
                val = ws.cell(1, c).value
                if isinstance(val, (datetime, date)):
                    vd = val.date() if isinstance(val, datetime) else val
                    if vd == target_date:
                        target_cols.append(c)
                elif isinstance(val, str):
                    p = parse_date_str(val)
                    if p and p == target_date:
                        target_cols.append(c)

            # If date column not identified, search all columns
            cols_to_check = target_cols if target_cols else range(1, ws.max_column + 1)

            for col_idx in cols_to_check:
                for row_idx in range(2, max(ws.max_row + 1, 50)):
                    cell = ws.cell(row_idx, col_idx)
                    cv = cell.value
                    if cv is None:
                        continue
                    cv_str = str(cv).strip()
                    # Match clean room or room with status suffix
                    if cv_str == clean_room or cv_str.startswith(f"{clean_room} ") or (clean_room.isdigit() and cv == int(clean_room)):
                        cell.value = new_cell_value
                        if fill_to_apply:
                            cell.fill = fill_to_apply
                        modified = True

        if modified:
            try:
                self._save_workbook(wb)
            except Exception as e:
                print(f"[BookingCallsManager] Error saving workbook after cell manipulation: {e}")

        return modified

    # -------------------------------------------------------------------------
    # Generate & Save Today's Calls JSON in 'BOOKING CALLS FOR TODAY'
    # -------------------------------------------------------------------------
    def generate_today_booking_calls_json(self, target_date: Optional[date] = None) -> List[Dict[str, Any]]:
        """
        Calculates today's scheduled calls strictly for BOOKING.COM,
        saving the output into DATABASE/BOOKING CALLS FOR TODAY/booking_calls_for_today.json.
        """
        if target_date is None:
            target_date = date.today()

        master_state = self.data_manager.load_master_state()
        calls_state = self.load_calls_state()
        date_key = target_date.strftime("%Y-%m-%d")
        saved_today = calls_state.get(date_key, {})

        scheduled_rooms_data: List[Dict[str, Any]] = []

        for b_id, booking in master_state.items():
            if not isinstance(booking, dict) or not is_booking_com(booking):
                continue

            room = str(booking.get("Δωμάτιο", "")).strip()
            arr_date = parse_date_str(booking.get("Άφιξη", ""))
            dep_date = parse_date_str(booking.get("Αναχώρηση", ""))

            if not room or not arr_date or not dep_date:
                continue

            sched_dates = calculate_call_schedule(arr_date, dep_date)
            if target_date in sched_dates:
                guest_names = booking.get("Πελάτες", [])
                saved_item = saved_today.get(room, {})

                call_record = {
                    "room": room,
                    "property": "SANDY BEACH",
                    "booking_id": b_id,
                    "guest_name": ", ".join(guest_names) if guest_names else "Guest",
                    "agency": "BOOKING.COM",
                    "arrival": booking.get("Άφιξη", ""),
                    "departure": booking.get("Αναχώρηση", ""),
                    "room_type": booking.get("Τύπος Δωματίου", ""),
                    "status": saved_item.get("status", "Green"),
                    "notes": saved_item.get("notes", ""),
                    "date": date_key
                }
                scheduled_rooms_data.append(call_record)

        self.save_today_calls_json(scheduled_rooms_data)
        return scheduled_rooms_data

    def save_room_call(self, room: str, status: str, notes: str,
                        call_date: Optional[date] = None,
                        additional_data: Optional[Dict[str, Any]] = None):
        """Saves status/notes to JSON and performs direct Excel cell manipulation."""
        if call_date is None:
            call_date = date.today()

        date_key = call_date.strftime("%Y-%m-%d")

        # 1. Update Master Calls State
        calls_state = self.load_calls_state()
        if date_key not in calls_state:
            calls_state[date_key] = {}

        entry = calls_state[date_key].get(room, {})
        entry["status"] = status
        entry["notes"] = notes
        entry["agency"] = "BOOKING.COM"
        entry["updated_at"] = datetime.now().isoformat()
        if additional_data:
            entry.update(additional_data)
        calls_state[date_key][room] = entry
        self.save_calls_state(calls_state)

        # 2. Update Today Calls JSON
        today_calls = self.load_today_calls_json()
        updated = False
        for c in today_calls:
            if c.get("room") == room:
                c["status"] = status
                c["notes"] = notes
                updated = True
                break
        if not updated:
            new_call = {
                "room": room,
                "status": status,
                "notes": notes,
                "agency": "BOOKING.COM",
                "date": date_key
            }
            if additional_data:
                new_call.update(additional_data)
            today_calls.append(new_call)

        self.save_today_calls_json(today_calls)

        # 3. Direct Excel Cell Manipulation
        self.update_excel_cell_status(room, status, call_date)


# =============================================================================
# Explicit Non-Wheel ComboBox
# =============================================================================

class ExplicitComboBox(QComboBox):
    def wheelEvent(self, e):
        e.ignore()


# =============================================================================
# PyQt6 Room Call Card Widget
# =============================================================================

class RoomCallCard(QFrame):
    status_changed = pyqtSignal(str, str, str)  # room, status, notes
    feedback_submitted = pyqtSignal(str, str)   # room, comment_text

    def __init__(self, call_data: Dict[str, Any], parent=None):
        super().__init__(parent)
        self.call_data = call_data
        self.room = str(call_data.get("room", ""))
        self.last_submitted_feedback = str(call_data.get("notes", "")).strip()
        self._init_ui()

    def _init_ui(self):
        self.setFrameShape(QFrame.Shape.StyledPanel)
        self.setObjectName("room_card")

        layout = QVBoxLayout(self)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(8)

        # Header Row
        top_row = QHBoxLayout()
        lbl_room = QLabel(f"🏠 Room {self.room}")
        lbl_room.setStyleSheet("font-size: 16px; font-weight: bold; color: #1E293B;")
        top_row.addWidget(lbl_room)
        top_row.addStretch()

        lbl_bcom = QLabel("BOOKING.COM")
        lbl_bcom.setStyleSheet("background-color: #003580; color: white; padding: 2px 6px; border-radius: 3px; font-weight: bold; font-size: 10px;")
        top_row.addWidget(lbl_bcom)

        lbl_prop = QLabel("SANDY BEACH")
        lbl_prop.setStyleSheet("background-color: #DBEAFE; color: #1E40AF; padding: 3px 8px; border-radius: 4px; font-weight: bold; font-size: 11px;")
        top_row.addWidget(lbl_prop)
        layout.addLayout(top_row)

        # Guest & Stay Info
        guest_text = self.call_data.get("guest_name", "Guest")
        b_id = self.call_data.get("booking_id", "")
        arr = self.call_data.get("arrival", "-")
        dep = self.call_data.get("departure", "-")

        lbl_guest = QLabel(f"👤 <b>{guest_text}</b>" + (f" (Booking #{b_id})" if b_id else ""))
        lbl_guest.setStyleSheet("color: #334155; font-size: 13px;")
        lbl_guest.setWordWrap(True)
        layout.addWidget(lbl_guest)

        lbl_stay = QLabel(f"📅 <b>Stay:</b> {arr} ➔ {dep}")
        lbl_stay.setStyleSheet("color: #64748B; font-size: 11px;")
        layout.addWidget(lbl_stay)

        # Status Selector
        status_row = QHBoxLayout()
        lbl_status_title = QLabel("Call Status:")
        lbl_status_title.setStyleSheet("font-weight: bold; color: #334155;")
        status_row.addWidget(lbl_status_title)

        self.cmb_status = ExplicitComboBox()
        self.cmb_status.addItems(CALL_STATUSES)
        self.cmb_status.setFocusPolicy(Qt.FocusPolicy.StrongFocus)

        current_status = self.call_data.get("status", "Green")
        if current_status in CALL_STATUSES:
            self.cmb_status.setCurrentText(current_status)
        else:
            self.cmb_status.setCurrentIndex(0)

        self.cmb_status.activated.connect(self._on_status_activated)
        status_row.addWidget(self.cmb_status, stretch=1)
        layout.addLayout(status_row)

        # Notes / Feedback
        notes_header = QHBoxLayout()
        lbl_notes = QLabel("Call Notes / Guest Feedback:")
        lbl_notes.setStyleSheet("font-weight: bold; color: #334155;")
        notes_header.addWidget(lbl_notes)
        notes_header.addStretch()

        self.lbl_feedback_badge = QLabel("")
        self.lbl_feedback_badge.setStyleSheet("color: #059669; font-size: 11px; font-weight: bold;")
        notes_header.addWidget(self.lbl_feedback_badge)
        layout.addLayout(notes_header)

        self.txt_notes = QTextEdit()
        self.txt_notes.setMaximumHeight(65)
        self.txt_notes.setPlaceholderText("Enter call details, guest feedback, or requests...")
        self.txt_notes.setText(self.call_data.get("notes", ""))
        self.txt_notes.setStyleSheet("""
            QTextEdit {
                background-color: #F8FAFC;
                border: 1px solid #CBD5E1;
                border-radius: 4px;
                padding: 6px;
                font-size: 12px;
            }
            QTextEdit:focus {
                border: 1px solid #003580;
                background-color: #FFFFFF;
            }
        """)
        layout.addWidget(self.txt_notes)

        # Action bar
        action_row = QHBoxLayout()
        self.btn_send_todo = QPushButton("📝 Add Feedback to To Do List")
        self.btn_send_todo.setStyleSheet("""
            QPushButton {
                background-color: #EFF6FF;
                border: 1px solid #3B82F6;
                color: #1D4ED8;
                font-weight: bold;
                font-size: 11px;
                padding: 4px 10px;
                border-radius: 4px;
            }
            QPushButton:hover {
                background-color: #3B82F6;
                color: white;
            }
        """)
        self.btn_send_todo.clicked.connect(self._submit_feedback_explicit)
        action_row.addStretch()
        action_row.addWidget(self.btn_send_todo)
        layout.addLayout(action_row)

        # Auto capture timer
        self.auto_capture_timer = QTimer(self)
        self.auto_capture_timer.setSingleShot(True)
        self.auto_capture_timer.setInterval(1500)
        self.auto_capture_timer.timeout.connect(self._auto_capture_feedback)

        self.txt_notes.textChanged.connect(self._on_notes_change)
        self._update_card_style(self.cmb_status.currentText())

    def _update_card_style(self, status: str):
        col = STATUS_COLORS.get(status, {"bg": "#FFFFFF", "border": "#CBD5E1", "text": "#000000"})
        self.setStyleSheet(f"""
            QFrame#room_card {{
                background-color: #FFFFFF;
                border: 2px solid {col['border']};
                border-radius: 8px;
            }}
            QComboBox {{
                background-color: {col['bg']};
                border: 1px solid {col['border']};
                color: {col['text']};
                font-weight: bold;
                padding: 4px 8px;
                border-radius: 4px;
            }}
        """)

    def _on_status_activated(self, index: int):
        new_status = self.cmb_status.itemText(index)
        self._update_card_style(new_status)
        self.status_changed.emit(self.room, new_status, self.txt_notes.toPlainText())

    def _on_notes_change(self):
        self.status_changed.emit(self.room, self.cmb_status.currentText(), self.txt_notes.toPlainText())
        self.auto_capture_timer.start()

    def _auto_capture_feedback(self):
        comment = self.txt_notes.toPlainText().strip()
        if comment and comment != self.last_submitted_feedback:
            self.last_submitted_feedback = comment
            self.feedback_submitted.emit(self.room, comment)
            self.lbl_feedback_badge.setText("✓ Added to To Do List")
            QTimer.singleShot(3000, lambda: self.lbl_feedback_badge.setText(""))

    def _submit_feedback_explicit(self):
        comment = self.txt_notes.toPlainText().strip()
        if comment:
            self.last_submitted_feedback = comment
            self.feedback_submitted.emit(self.room, comment)
            self.lbl_feedback_badge.setText("✓ Added to To Do List")
            QTimer.singleShot(3000, lambda: self.lbl_feedback_badge.setText(""))

    def get_data(self) -> Dict[str, Any]:
        return {
            "room": self.room,
            "property": "SANDY BEACH",
            "booking_id": self.call_data.get("booking_id", ""),
            "guest_name": self.call_data.get("guest_name", ""),
            "agency": "BOOKING.COM",
            "arrival": self.call_data.get("arrival", ""),
            "departure": self.call_data.get("departure", ""),
            "status": self.cmb_status.currentText(),
            "notes": self.txt_notes.toPlainText()
        }


# =============================================================================
# PyQt6 BOOKING CALLS Main Page Widget
# =============================================================================

class BookingCallsWidget(QWidget):
    """
    Main CRM interface for the BOOKING CALLS module in app.py.
    Strictly displays today's scheduled calls from:
    'DATABASE/BOOKING CALLS FOR TODAY/booking_calls_for_today.json'.
    """

    feedback_submitted = pyqtSignal(str, str)  # room, comment_text

    def __init__(self, parent=None):
        super().__init__(parent)
        self.manager = BookingCallsManager()
        self.cards: List[RoomCallCard] = []
        self._init_ui()
        self.refresh_calls()

    def _init_ui(self):
        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(16, 16, 16, 16)
        main_layout.setSpacing(12)

        # Header Banner
        header_card = QFrame()
        header_card.setStyleSheet("""
            QFrame {
                background: qlineargradient(x1:0, y1:0, x2:1, y2:0, stop:0 #003580, stop:1 #006CE4);
                border-radius: 8px;
                padding: 14px;
            }
            QLabel { color: white; }
        """)
        h_layout = QHBoxLayout(header_card)

        v_titles = QVBoxLayout()
        lbl_title = QLabel("📞 BOOKING CALLS CRM (BOOKING.COM ONLY)")
        lbl_title.setFont(QFont("Segoe UI", 16, QFont.Weight.Bold))
        today_str = date.today().strftime("%A, %d %B %Y")
        lbl_sub = QLabel(f"Scheduled follow-up calls for today ({today_str}) — Source: DATABASE/BOOKING CALLS FOR TODAY/")
        lbl_sub.setStyleSheet("color: #DBEAFE; font-size: 12px;")
        v_titles.addWidget(lbl_title)
        v_titles.addWidget(lbl_sub)
        h_layout.addLayout(v_titles)
        h_layout.addStretch()

        self.lbl_total_badge = QLabel("Booking.com Calls: 0")
        self.lbl_total_badge.setStyleSheet(
            "background-color: rgba(255, 255, 255, 0.2); padding: 6px 12px; border-radius: 6px; font-weight: bold;"
        )
        h_layout.addWidget(self.lbl_total_badge)
        main_layout.addWidget(header_card)

        # Control Bar
        ctrl_bar = QHBoxLayout()

        self.txt_search = QLineEdit()
        self.txt_search.setPlaceholderText("🔍 Search Room, Guest, or Booking ID...")
        self.txt_search.setFixedWidth(280)
        self.txt_search.textChanged.connect(self._apply_filters)
        ctrl_bar.addWidget(self.txt_search)

        self.cmb_filter_status = QComboBox()
        self.cmb_filter_status.addItems(["All Statuses"] + CALL_STATUSES)
        self.cmb_filter_status.currentTextChanged.connect(self._apply_filters)
        ctrl_bar.addWidget(self.cmb_filter_status)

        ctrl_bar.addStretch()

        btn_sync = QPushButton("🔄 Sync Excel & Today's JSON")
        btn_sync.setStyleSheet("""
            QPushButton {
                background-color: #10B981;
                color: white;
                font-weight: bold;
                padding: 7px 14px;
                border-radius: 4px;
                border: none;
            }
            QPushButton:hover { background-color: #059669; }
        """)
        btn_sync.clicked.connect(self.sync_and_refresh)
        ctrl_bar.addWidget(btn_sync)

        btn_save_all = QPushButton("💾 Save All Changes")
        btn_save_all.setStyleSheet("""
            QPushButton {
                background-color: #003580;
                color: white;
                font-weight: bold;
                padding: 7px 14px;
                border-radius: 4px;
                border: none;
            }
            QPushButton:hover { background-color: #00224F; }
        """)
        btn_save_all.clicked.connect(self.save_all)
        ctrl_bar.addWidget(btn_save_all)

        main_layout.addLayout(ctrl_bar)

        # Scrollable Cards Grid Container
        self.scroll_area = QScrollArea()
        self.scroll_area.setWidgetResizable(True)
        self.scroll_area.setStyleSheet("QScrollArea { border: none; background-color: transparent; }")

        self.cards_container = QWidget()
        self.cards_container.setObjectName("cards_container")
        self.cards_container.setStyleSheet("QWidget#cards_container { background-color: transparent; }")
        self.grid_layout = QGridLayout(self.cards_container)
        self.grid_layout.setSpacing(14)
        self.grid_layout.setAlignment(Qt.AlignmentFlag.AlignTop)

        self.scroll_area.setWidget(self.cards_container)
        main_layout.addWidget(self.scroll_area, stretch=1)

    def refresh_calls(self):
        """Strictly loads today's calls from DATABASE/BOOKING CALLS FOR TODAY/booking_calls_for_today.json."""
        for card in self.cards:
            self.grid_layout.removeWidget(card)
            card.deleteLater()
        self.cards.clear()

        calls_data = self.manager.load_today_calls_json()
        self.lbl_total_badge.setText(f"Booking.com Scheduled: {len(calls_data)}")

        if not calls_data:
            lbl_empty = QLabel(
                "No Booking.com calls scheduled for today in 'DATABASE/BOOKING CALLS FOR TODAY/'.\n"
                "Click '🔄 Sync Excel & Today\\'s JSON' to calculate the schedule."
            )
            lbl_empty.setStyleSheet("color: #64748B; font-size: 14px; font-style: italic; padding: 20px;")
            self.grid_layout.addWidget(lbl_empty, 0, 0)
            return

        cols = 2
        for i, c_data in enumerate(calls_data):
            row = i // cols
            col = i % cols
            card = RoomCallCard(c_data)
            card.status_changed.connect(self._on_card_status_changed)
            card.feedback_submitted.connect(self.feedback_submitted.emit)
            self.grid_layout.addWidget(card, row, col)
            self.cards.append(card)

    def _on_card_status_changed(self, room: str, status: str, notes: str):
        self.manager.save_room_call(room, status, notes)

    def _apply_filters(self):
        search_query = self.txt_search.text().strip().lower()
        status_filter = self.cmb_filter_status.currentText()

        visible_count = 0
        for card in self.cards:
            data = card.get_data()
            room = data["room"].lower()
            guest = data["guest_name"].lower()
            b_id = data["booking_id"].lower()
            status = data["status"]

            match_search = (
                not search_query
                or search_query in room
                or search_query in guest
                or search_query in b_id
            )

            match_status = (
                status_filter == "All Statuses"
                or status == status_filter
            )

            is_visible = match_search and match_status
            card.setVisible(is_visible)
            if is_visible:
                visible_count += 1

        self.lbl_total_badge.setText(f"Showing: {visible_count} of {len(self.cards)}")

    def sync_and_refresh(self):
        """Syncs Sheet 1 (Arrivals) and Sheet 2 (Follow-Up) for Booking.com only and updates today's JSON."""
        try:
            arr_rooms = self.manager.sync_today_arrivals_to_excel()
            schedule = self.manager.sync_followup_schedule_to_excel()
            today_calls = self.manager.generate_today_booking_calls_json()

            self.refresh_calls()
            QMessageBox.information(
                self,
                "Sync Complete",
                f"Booking.com Calls synchronized successfully!\n\n"
                f"• Sheet 1 (ARRIVALS): {len(arr_rooms)} Booking.com arrivals populated\n"
                f"• Sheet 2 (FOLLOW UP): {len(schedule)} follow-up dates scheduled\n"
                f"• Today's Scheduled Calls: {len(today_calls)} rooms\n\n"
                f"Output saved strictly to: 'DATABASE/BOOKING CALLS FOR TODAY/booking_calls_for_today.json'"
            )
        except Exception as e:
            QMessageBox.critical(self, "Sync Error", f"Failed to synchronize Booking Calls:\n{str(e)}")

    def save_all(self):
        """Saves all room statuses and notes from all cards."""
        for card in self.cards:
            data = card.get_data()
            self.manager.save_room_call(
                room=data["room"],
                status=data["status"],
                notes=data["notes"],
                additional_data=data
            )
        QMessageBox.information(self, "Saved", "All Booking.com call records have been saved successfully!")
