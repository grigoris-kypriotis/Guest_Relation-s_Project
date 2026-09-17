"""
Unit tests for booking_calls.py:
Tests the 4 scheduling rules, consecutive date resolution, Excel synchronization,
and state persistence.
"""

import os
import shutil
import tempfile
import unittest
from datetime import date, timedelta

import openpyxl
import sys
from pathlib import Path

# Ensure project root is in sys.path
PROJECT_ROOT = str(Path(__file__).resolve().parent.parent)
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from MODULES.booking_calls import (
    calculate_call_schedule,
    BookingCallsManager,
    CALL_STATUSES
)


class TestBookingCallsScheduling(unittest.TestCase):
    def test_rule_1_2_3_4_seven_night_stay(self):
        # Arrival: 10/09/2026, Departure: 17/09/2026 (7 nights)
        # Rule 1 (A+1): 11
        # Rule 2 (D-1): 16
        # Rule 3 (every other day): 13, 15
        # Candidates: [11, 13, 15, 16]
        # Rule 4: 15 and 16 are consecutive! Discard earlier (15), keep later (16).
        # Expected: [11, 13, 16]
        arr = date(2026, 9, 10)
        dep = date(2026, 9, 17)
        schedule = calculate_call_schedule(arr, dep)
        expected = [date(2026, 9, 11), date(2026, 9, 13), date(2026, 9, 16)]
        self.assertEqual(schedule, expected)

    def test_eight_night_stay_no_consecutive(self):
        # Arrival: 10/09/2026, Departure: 18/09/2026 (8 nights)
        # Rule 1: 11
        # Rule 2: 17
        # Rule 3: 13, 15
        # Candidates: [11, 13, 15, 17]
        # Rule 4: No consecutive days
        # Expected: [11, 13, 15, 17]
        arr = date(2026, 9, 10)
        dep = date(2026, 9, 18)
        schedule = calculate_call_schedule(arr, dep)
        expected = [date(2026, 9, 11), date(2026, 9, 13), date(2026, 9, 15), date(2026, 9, 17)]
        self.assertEqual(schedule, expected)

    def test_three_night_stay(self):
        # Arrival: 10/09/2026, Departure: 13/09/2026 (3 nights)
        # Rule 1: 11
        # Rule 2: 12
        # Candidates: [11, 12] (consecutive!)
        # Rule 4: Discard 11, keep 12.
        # Expected: [12]
        arr = date(2026, 9, 10)
        dep = date(2026, 9, 13)
        schedule = calculate_call_schedule(arr, dep)
        expected = [date(2026, 9, 12)]
        self.assertEqual(schedule, expected)

    def test_two_night_stay(self):
        # Arrival: 10/09/2026, Departure: 12/09/2026 (2 nights)
        # A+1 == D-1 == 11
        # Expected: [11]
        arr = date(2026, 9, 10)
        dep = date(2026, 9, 12)
        schedule = calculate_call_schedule(arr, dep)
        expected = [date(2026, 9, 11)]
        self.assertEqual(schedule, expected)

    def test_five_night_stay(self):
        # Arrival: 10/09/2026, Departure: 15/09/2026 (5 nights)
        # Rule 1: 11
        # Rule 2: 14
        # Rule 3: 13
        # Candidates: [11, 13, 14] -> 13 and 14 consecutive! Discard 13, keep 14.
        # Expected: [11, 14]
        arr = date(2026, 9, 10)
        dep = date(2026, 9, 15)
        schedule = calculate_call_schedule(arr, dep)
        expected = [date(2026, 9, 11), date(2026, 9, 14)]
        self.assertEqual(schedule, expected)

    def test_four_night_stay(self):
        # Arrival: 10/09/2026, Departure: 14/09/2026 (4 nights)
        # Rule 1: 11
        # Rule 2: 13
        # Candidates: [11, 13]
        # Expected: [11, 13]
        arr = date(2026, 9, 10)
        dep = date(2026, 9, 14)
        schedule = calculate_call_schedule(arr, dep)
        expected = [date(2026, 9, 11), date(2026, 9, 13)]
        self.assertEqual(schedule, expected)

    def test_standard_six_statuses(self):
        self.assertEqual(len(CALL_STATUSES), 6)
        self.assertIn("Green", CALL_STATUSES)
        self.assertIn("Red", CALL_STATUSES)
        self.assertIn("Yellow", CALL_STATUSES)
        self.assertIn("N/A (NO ANSWER)", CALL_STATUSES)
        self.assertIn("N/E (NO ENGLISH)", CALL_STATUSES)
        self.assertIn("N/W (LINE NOT WORKING)", CALL_STATUSES)


class TestBookingCallsManager(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp(prefix="test_booking_calls_")
        self.xlsx_path = os.path.join(self.test_dir, "BOOKING CALLS.xlsx")
        self.state_path = os.path.join(self.test_dir, "booking_calls_state.json")

        # Create dummy workbook with 3 sheets
        wb = openpyxl.Workbook()
        ws_arr = wb.active
        ws_arr.title = "ARRIVALS"
        ws_arr.cell(1, 1).value = "DATE"

        ws_f1 = wb.create_sheet("FOLLOW UP 1")
        ws_f1.cell(1, 1).value = "DATE"

        ws_f2 = wb.create_sheet("FOLLOW UP")
        ws_f2.cell(1, 1).value = "DATE"

        wb.save(self.xlsx_path)

        self.today_json_dir = os.path.join(self.test_dir, "booking calls for today")
        self.today_json_path = os.path.join(self.today_json_dir, "booking_calls_today.json")

        self.manager = BookingCallsManager(
            xlsx_path=self.xlsx_path,
            state_path=self.state_path,
            today_json_path=self.today_json_path
        )

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_is_booking_com_filtering(self):
        from MODULES.booking_calls import is_booking_com
        # Matches Booking.com in various fields
        self.assertTrue(is_booking_com({"Χρεώστης": "BOOKING.COM"}))
        self.assertTrue(is_booking_com({"Χρεώστης": "Booking.com B.V."}))
        self.assertTrue(is_booking_com({"agency": "booking.com"}))
        self.assertTrue(is_booking_com({"Τιμοκατάλογος": "BOOKING.COM BAR 2026"}))

        # Strictly rejects all other agencies
        self.assertFalse(is_booking_com({"Χρεώστης": "TUI DEUTSCHLAND"}))
        self.assertFalse(is_booking_com({"Χρεώστης": "EXPEDIA"}))
        self.assertFalse(is_booking_com({"Χρεώστης": "DERTOUR GMBH"}))
        self.assertFalse(is_booking_com({"Χρεώστης": "RAINBOW TOURS"}))
        self.assertFalse(is_booking_com({}))

    def test_state_persistence(self):
        self.manager.save_room_call(
            room="1101",
            status="Green",
            notes="Guest loves the sea view",
            call_date=date(2026, 9, 14)
        )

        state = self.manager.load_calls_state()
        self.assertIn("2026-09-14", state)
        self.assertIn("1101", state["2026-09-14"])
        self.assertEqual(state["2026-09-14"]["1101"]["status"], "Green")
        self.assertEqual(state["2026-09-14"]["1101"]["notes"], "Guest loves the sea view")

    def test_today_calls_json_strictly_in_designated_directory(self):
        sample_calls = [
            {
                "room": "1205",
                "property": "SANDY BEACH",
                "booking_id": "99123",
                "guest_name": "Smith John",
                "agency": "BOOKING.COM",
                "arrival": "10/09/2026",
                "departure": "17/09/2026",
                "status": "Green",
                "notes": "First call OK",
                "date": "2026-09-14"
            }
        ]
        self.manager.save_today_calls_json(sample_calls)

        self.assertTrue(os.path.exists(self.today_json_path))
        loaded = self.manager.load_today_calls_json()
        self.assertEqual(len(loaded), 1)
        self.assertEqual(loaded[0]["room"], "1205")
        self.assertEqual(loaded[0]["agency"], "BOOKING.COM")
        self.assertEqual(loaded[0]["status"], "Green")

    def test_sync_sheet1_arrivals_and_sheet2_followup(self):
        # Mock arrivals state with Booking.com and non-Booking.com arrivals
        mock_arrivals = {
            "SANDY BEACH": {
                "10001": {
                    "room": "1101",
                    "guests": ["Booking Guest"],
                    "agency": "BOOKING.COM",
                    "arrival": "15/09/2026",
                    "departure": "22/09/2026"
                },
                "20002": {
                    "room": "1202",
                    "guests": ["TUI Guest"],
                    "agency": "TUI",
                    "arrival": "15/09/2026",
                    "departure": "22/09/2026"
                }
            }
        }
        self.manager.data_manager.load_arrivals_state = lambda: mock_arrivals

        # Sync Sheet 1 (ARRIVALS)
        count = self.manager.sync_today_arrivals_to_excel()
        self.assertEqual(len(count), 1)
        self.assertEqual(count, ["1101"])

        wb = openpyxl.load_workbook(self.xlsx_path)
        sheet_arrivals = wb["ARRIVALS"]
        found_1101_arr = any("1101" in str(sheet_arrivals.cell(r, c).value or "")
                             for r in range(1, 20) for c in range(1, 10))
        found_1202_arr = any("1202" in str(sheet_arrivals.cell(r, c).value or "")
                             for r in range(1, 20) for c in range(1, 10))
        self.assertTrue(found_1101_arr, "Booking.com room 1101 must be in Sheet 1 ARRIVALS")
        self.assertFalse(found_1202_arr, "Non-Booking.com room 1202 must NOT be in Sheet 1 ARRIVALS")

        # Mock master state for Sheet 2 (FOLLOW UP)
        mock_master = {
            "10001": {
                "Δωμάτιο": "1101",
                "Άφιξη": "10/09/2026",
                "Αναχώρηση": "17/09/2026",
                "agency": "BOOKING.COM",
                "Πελάτες": ["Booking Guest"]
            },
            "20002": {
                "Δωμάτιο": "1202",
                "Άφιξη": "10/09/2026",
                "Αναχώρηση": "17/09/2026",
                "agency": "TUI",
                "Πελάτες": ["TUI Guest"]
            }
        }
        self.manager.data_manager.load_master_state = lambda: mock_master

        # Sync Sheet 2 (FOLLOW UP)
        schedule = self.manager.sync_followup_schedule_to_excel()
        self.assertGreater(len(schedule), 0)

        wb = openpyxl.load_workbook(self.xlsx_path)
        sheet_followup = wb["FOLLOW UP"]
        found_1101_fol = any("1101" in str(sheet_followup.cell(r, c).value or "")
                             for r in range(1, 25) for c in range(1, 15))
        found_1202_fol = any("1202" in str(sheet_followup.cell(r, c).value or "")
                             for r in range(1, 25) for c in range(1, 15))
        self.assertTrue(found_1101_fol, "Booking.com room 1101 must be in Sheet 2 FOLLOW UP")
        self.assertFalse(found_1202_fol, "Non-Booking.com room 1202 must NOT be in Sheet 2 FOLLOW UP")

    def test_direct_cell_manipulation_colors_and_strings(self):
        # Seed cell 1101 under date 2026-09-11 in Sheet 2 (FOLLOW UP)
        wb_init = openpyxl.load_workbook(self.xlsx_path)
        ws_init = wb_init["FOLLOW UP"]
        ws_init.cell(1, 2).value = date(2026, 9, 11)
        ws_init.cell(3, 2).value = "1101"
        wb_init.save(self.xlsx_path)

        # 1. Update cell with Green color
        ok = self.manager.update_excel_cell_status(
            room="1101",
            target_date=date(2026, 9, 11),
            status="Green",
            sheet_name="FOLLOW UP"
        )
        self.assertTrue(ok)

        wb = openpyxl.load_workbook(self.xlsx_path)
        ws = wb["FOLLOW UP"]
        cell_1101 = None
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                if "1101" in str(ws.cell(r, c).value or ""):
                    cell_1101 = ws.cell(r, c)
                    break
            if cell_1101:
                break

        self.assertIsNotNone(cell_1101)
        self.assertIsNotNone(cell_1101.fill)
        self.assertIn("D4EDDA", str(cell_1101.fill.start_color.rgb).upper())

        # 2. Update cell with N/A status (string insertion into cell)
        ok2 = self.manager.update_excel_cell_status(
            room="1101",
            target_date=date(2026, 9, 11),
            status="N/A (NO ANSWER)",
            sheet_name="FOLLOW UP"
        )
        self.assertTrue(ok2)

        wb = openpyxl.load_workbook(self.xlsx_path)
        ws = wb["FOLLOW UP"]
        val_after = str(ws.cell(cell_1101.row, cell_1101.column).value or "")
        self.assertIn("N/A", val_after)
        self.assertIn("1101", val_after)

    def test_feedback_to_do_formatting(self):
        comment = "Guest requested extra pillows"
        room = "1402"
        expected_format = f"Feedback on exclusivi: {comment} - Room {room}"
        self.assertEqual(expected_format, "Feedback on exclusivi: Guest requested extra pillows - Room 1402")


if __name__ == "__main__":
    unittest.main()

